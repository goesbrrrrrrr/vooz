# -*- coding: utf-8 -*-
"""Вспомогательный модуль: оформление отчёта ВГУ в стиле отчётов ЛР1-ЛР3."""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def new_doc():
    doc = Document()
    doc._headings = []          # (уровень, текст) — для автосборки содержания
    doc._toc_anchor = None
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(3)
        s.right_margin = Cm(1.5)

    for name, size, bold in (("Heading 1", 16, True), ("Heading 2", 14, True),
                             ("Heading 3", 14, True)):
        h = doc.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = bold
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.line_spacing = 1.5
    return doc


def _p(doc, text="", *, align=None, bold=False, italic=False, size=None,
       indent=None, spacing=None, style=None, font=None):
    par = doc.add_paragraph(style=style)
    if align is not None:
        par.alignment = align
    par.paragraph_format.first_line_indent = Cm(0 if indent is None else indent)
    if spacing is not None:
        par.paragraph_format.line_spacing = spacing
    if text:
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.name = font or FONT
        if size:
            r.font.size = Pt(size)
    return par


def title_page(doc, lab_no, theme_lines):
    C = WD_ALIGN_PARAGRAPH.CENTER
    for t in ("МИНОБРНАУКИ РОССИИ",
              "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ ОБРАЗОВАТЕЛЬНОЕ",
              "УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
              "«ВОРОНЕЖСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ»",
              "(ФГБОУ ВО «ВГУ»)"):
        _p(doc, t, align=C, size=13, spacing=1.0)
    _p(doc, "", spacing=1.0)
    _p(doc, "Факультет прикладной математики, информатики и механики",
       align=C, size=13, spacing=1.0)
    for _ in range(4):
        _p(doc, "", spacing=1.0)
    _p(doc, "Дисциплина: «Численные методы»", align=C, spacing=1.0)
    _p(doc, "Лабораторное задание № %d" % lab_no, align=C, spacing=1.0)
    for _ in range(4):
        _p(doc, "", spacing=1.0)
    _p(doc, "Отчет по лабораторной работе № %d" % lab_no,
       align=C, bold=True, size=16, spacing=1.0)
    for i, t in enumerate(theme_lines):
        _p(doc, ("Тема: " if i == 0 else "") + t, align=C, bold=True,
           size=16, spacing=1.0)
    for _ in range(5):
        _p(doc, "", spacing=1.0)
    R = WD_ALIGN_PARAGRAPH.RIGHT
    _p(doc, "Выполнил:", align=R, spacing=1.0)
    _p(doc, "студент 3 курса 8 группы", align=R, spacing=1.0)
    _p(doc, "Юрканов Артем Романович", align=R, spacing=1.0)
    _p(doc, "", spacing=1.0)
    _p(doc, "Проверил:", align=R, spacing=1.0)
    _p(doc, "преподаватель Махинова О.А.", align=R, spacing=1.0)
    for _ in range(3):
        _p(doc, "", spacing=1.0)
    _p(doc, "Воронеж 2026", align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def contents_placeholder(doc):
    """
    Резервирует место под содержание. Реальные строки вставляются позже,
    функцией fill_contents, когда станут известны номера страниц.
    """
    doc._toc_anchor = doc.add_paragraph()
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc._toc_anchor


def _toc_line(par, lvl, txt, page):
    """Одна строка содержания: текст, точки-выноска, номер страницы."""
    par.paragraph_format.first_line_indent = Cm(1.0 if lvl else 0)
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # позиция табуляции с точечным заполнителем у правого поля
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9355")
    tabs.append(tab)
    par._p.get_or_add_pPr().append(tabs)
    r = par.add_run(txt)
    r.bold = (lvl == 0)
    r.font.name = FONT
    r.font.size = Pt(13 if lvl else 14)
    r2 = par.add_run("\t%s" % page)
    r2.bold = (lvl == 0)
    r2.font.name = FONT
    r2.font.size = Pt(13 if lvl else 14)


def fill_contents(doc, pages=None):
    """
    Заполняет содержание по фактически добавленным заголовкам.
    pages — словарь {текст заголовка: номер страницы}; если его нет,
    ставятся прочерки (первый проход, до вёрстки).
    """
    anchor = doc._toc_anchor
    head = anchor.insert_paragraph_before()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.paragraph_format.first_line_indent = Cm(0)
    r = head.add_run("Содержание")
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    anchor.insert_paragraph_before()
    for lvl, txt in doc._headings:
        _toc_line(anchor.insert_paragraph_before(), lvl, txt,
                  (pages or {}).get(txt, "—"))
    # убираем сам якорь
    anchor._p.getparent().remove(anchor._p)


def h1(doc, text):
    doc._headings.append((0, text))
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc._headings.append((1, text))
    doc.add_heading(text, level=2)


def para(doc, text):
    _p(doc, text, indent=1.25)


def formula(doc, text):
    _p(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True,
       spacing=1.0, font="Cambria Math")
    _p(doc, "", spacing=1.0)


def mono(doc, lines, size=8.5):
    """Блок моноширинного текста (листинг, вывод программы)."""
    for ln in lines:
        p = _p(doc, spacing=1.0)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(14)


def _fresh_num_id(doc):
    """
    Новый экземпляр нумерации, начинающийся с единицы.

    Стиль «List Number» в шаблоне ссылается на один-единственный numId,
    поэтому все нумерованные списки документа делят общий счётчик: второй
    список продолжается с того номера, на котором кончился первый. Чтобы
    каждый список начинался с 1, заводим отдельный w:num на том же
    abstractNum и переопределяем в нём стартовое значение.
    """
    numbering = doc.part.numbering_part.element
    # abstractNum, который использует стиль
    style_pPr = doc.styles["List Number"].element.find(qn("w:pPr"))
    src_id = style_pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    abstract = None
    for n in numbering.findall(qn("w:num")):
        if n.get(qn("w:numId")) == src_id:
            abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
            break
    used = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = str(max(used) + 1)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), new_id)
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract)
    num.append(ref)
    ov = OxmlElement("w:lvlOverride")
    ov.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    ov.append(start)
    num.append(ov)
    numbering.append(num)
    return new_id


def numbered(doc, items):
    num_id = _fresh_num_id(doc)          # у каждого списка свой счётчик
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        numPr = p._p.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0
        numPr.get_or_add_numId().val = int(num_id)
        r = p.add_run(it)
        r.font.name = FONT
        r.font.size = Pt(14)


def caption(doc, text):
    _p(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12,
       spacing=1.0)
    _p(doc, "", spacing=1.0)


def table(doc, header, rows, widths=None, size=11):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(htxt)
        r.bold = True
        r.font.name = FONT
        r.font.size = Pt(size)
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), "E8E8E8")
        hdr[i]._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            r = p.add_run(str(v))
            r.font.name = FONT
            r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    _p(doc, "", spacing=1.0)
    return t


def picture(doc, path, width_cm=16.0):
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(path, width=Cm(width_cm))


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_numbers(doc):
    """Номера страниц в нижнем колонтитуле по центру."""
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        for el, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                                ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
                                ("w:fldChar", {"w:fldCharType": "end"}, None)):
            e = OxmlElement(el)
            for k, v in attrs.items():
                e.set(qn(k), v)
            if text:
                e.text = text
            run._r.append(e)
        run.font.name = FONT
        run.font.size = Pt(12)
